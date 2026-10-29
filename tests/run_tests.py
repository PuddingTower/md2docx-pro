import os
import sys
import unittest

script_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if script_dir not in sys.path:
    sys.path.insert(0, script_dir)

from core.cleaner import MarkdownCleaner, CodeProtector
from core.templates import TemplateManager
from core.converter import DocumentConverter
from core.postprocessor import DocxPostProcessor
import docx
from docx.oxml.ns import qn


class TestMdToDocx(unittest.TestCase):

    def setUp(self):
        self.cleaner = MarkdownCleaner()
        self.converter = DocumentConverter()
        self.test_md = os.path.join(script_dir, "tests", "test_sample.md")

    def test_cleaner_structural_spacing(self):
        # Heading without blank line
        dirty_text = "正文紧贴下一行\n# 标题一"
        fixed = self.cleaner.fix_markdown_structural_spacing(dirty_text)
        self.assertEqual(fixed, "正文紧贴下一行\n\n# 标题一")

        # Table without blank line
        dirty_table = "正文紧贴表格\n| A | B |\n|---|---|"
        fixed_t = self.cleaner.fix_markdown_structural_spacing(dirty_table)
        self.assertEqual(fixed_t, "正文紧贴表格\n\n| A | B |\n|---|---|")

    def test_cleaner_math_latex(self):
        # Multiline display math without aligned
        math_text = "$$ a = 1 \\\\ b = 2 $$"
        cleaned = self.cleaner.clean_latex_syntax(math_text)
        self.assertIn(r"\begin{aligned}", cleaned)

    def test_cleaner_preserves_cjk_punctuation(self):
        raw = "标题：宽表，测试（括号）以及全角数字１２３"
        cleaned = self.cleaner.global_text_clean(raw)
        self.assertEqual(cleaned, "标题：宽表，测试（括号）以及全角数字123")

    def test_table_layout_full_width_and_header_repeat(self):
        res = self.converter.convert_file(self.test_md, preset="modern")
        self.assertTrue(res.success)
        doc = docx.Document(res.output_path)
        table = doc.tables[0]
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        self.assertEqual(tblW.attrib.get(qn('w:type')), 'dxa')
        self.assertEqual(int(tblW.attrib.get(qn('w:w'))), 9412)
        grid = [int(gc.get(qn('w:w'))) for gc in table._tbl.tblGrid.gridCol_lst]
        self.assertEqual(sum(grid), 9412)
        # Header repeats on page breaks; rows never split
        trPr = table.rows[0]._tr.trPr
        self.assertIsNotNone(trPr.find(qn('w:tblHeader')))
        self.assertIsNotNone(table.rows[1]._tr.trPr.find(qn('w:cantSplit')))

    def test_table_caption_style(self):
        res = self.converter.convert_file(self.test_md, preset="academic")
        self.assertTrue(res.success)
        doc = docx.Document(res.output_path)
        captions = [p for p in doc.paragraphs if p.style.name == "Table Caption"]
        self.assertTrue(captions)
        self.assertIn("表 1", captions[0].text)

    def test_conversion_in_same_directory(self):
        # Convert test_sample.md
        res = self.converter.convert_file(self.test_md, preset="modern")
        self.assertTrue(res.success)
        expected_docx = os.path.splitext(self.test_md)[0] + ".docx"
        self.assertEqual(os.path.abspath(res.output_path), os.path.abspath(expected_docx))
        self.assertTrue(os.path.exists(res.output_path))
        self.assertGreater(res.file_size_bytes, 1000)

        # Inspect generated docx
        doc = docx.Document(res.output_path)
        self.assertGreater(len(doc.tables), 0)
        table = doc.tables[0]
        # Verify cell padding and header
        cell_0 = table.rows[0].cells[0]
        tcPr = cell_0._tc.tcPr
        shd = tcPr.find(qn('w:shd'))
        self.assertIsNotNone(shd)
        self.assertEqual(shd.attrib.get(qn('w:fill')), "F1F5F9")

    def test_clipboard_text_conversion(self):
        text_content = """# 剪贴板测试文档
这是从剪贴板粘贴的一段纯文本。

| 项目 | 数值 |
| :--- | :--- |
| A | 100 |
"""
        temp_dir = os.path.join(script_dir, "tests")
        res = self.converter.convert_text(text_content, preset="academic", output_dir=temp_dir)
        self.assertTrue(res.success)
        self.assertTrue(os.path.exists(res.output_path))
        self.assertIn("剪贴板测试文档.docx", res.output_path)

        # Academic table inspection (Three-line table)
        doc = docx.Document(res.output_path)
        table = doc.tables[0]
        # header top border size should be 12 (1.5pt)
        tcPr = table.rows[0].cells[0]._tc.tcPr
        tcBorders = tcPr.find(qn('w:tcBorders'))
        self.assertIsNotNone(tcBorders)
        top_b = tcBorders.find(qn('w:top'))
        self.assertEqual(top_b.attrib.get(qn('w:sz')), '12')

    def test_gui_initialization_headless(self):
        from PyQt6.QtWidgets import QApplication
        from ui.main_window import MainWindow

        app = QApplication.instance()
        if app is None:
            app = QApplication(["--platform", "offscreen"])
        
        window = MainWindow()
        self.assertIsNotNone(window.drop_zone)
        self.assertIsNotNone(window.tray)
        self.assertEqual(window.combo_preset.count(), 3)
        window.close()


if __name__ == "__main__":
    unittest.main()
