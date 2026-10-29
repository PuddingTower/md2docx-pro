import os
import re
import logging
import unicodedata
from typing import List, Optional, Tuple
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger("md2docx.postprocessor")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Content width of the generated page (A4 minus 2.2cm side margins), in twips.
CONTENT_W = 11906 - 2 * 1247

TABLE_DESIGN = {
    "modern": dict(
        latin="Segoe UI", ea="Microsoft YaHei", size=10.5, color="1F2937",
        header_fill="F1F5F9", header_color="0F172A", zebra="F7F9FC",
        header_top=None, header_bottom=("single", 8, "CBD5E1"),
        row_line=("single", 4, "E7EDF5"), last_bottom=("single", 8, "94A3B8"),
        pad=(70, 70, 110, 110),
    ),
    "academic": dict(
        latin="Times New Roman", ea="SimSun", size=10.5, color="000000",
        header_fill=None, header_color="000000", zebra=None,
        header_top=("single", 12, "000000"), header_bottom=("single", 6, "000000"),
        row_line=None, last_bottom=("single", 12, "000000"),
        pad=(60, 60, 90, 90),
    ),
    "minimal": dict(
        latin="Arial", ea="Microsoft YaHei", size=10, color="1F2937",
        header_fill=None, header_color="111827", zebra=None,
        header_top=("single", 8, "111827"), header_bottom=("single", 6, "9CA3AF"),
        row_line=None, last_bottom=("single", 8, "111827"),
        pad=(60, 60, 100, 100),
    ),
}

CAPTION_DESIGN = {
    "modern": dict(size=9.5, color="64748B", bold=False),
    "academic": dict(size=10.5, color="000000", bold=True),
    "minimal": dict(size=9, color="6B7280", bold=False),
}

TABLE_CAPTION_RE = re.compile(r"^\s*(表|Table|Tab\.)\s*[0-9一二三四五六七八九十]+([\-–—.．]\s*\d+)*")
FIG_CAPTION_RE = re.compile(r"^\s*(图|Figure|Fig\.?)\s*[0-9一二三四五六七八九十]+([\-–—.．]\s*\d+)?")


def _display_width(text: str) -> int:
    width = 0
    for ch in text or "":
        width += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return width


MATH_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"


def _cell_units(cell) -> int:
    """Widest line of the cell, counting native equations (m:t) that p.text ignores."""
    widest = 0
    for p in cell.paragraphs:
        units = _display_width(p.text)
        for math in p._p.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"):
            inner = sum(_display_width(node.text) for node in math.iter(MATH_T))
            units += int(inner * 1.6) + 2
        widest = max(widest, units)
    return widest


def _insert_ordered(parent, node, order: List[str]):
    tag = node.tag.split('}')[-1]
    idx = order.index(tag) if tag in order else len(order)
    pos = 0
    for i, sibling in enumerate(list(parent)):
        sib_tag = sibling.tag.split('}')[-1]
        if sib_tag in order and order.index(sib_tag) > idx:
            break
        pos = i + 1
    parent.insert(pos, node)


TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
    "cellDel", "cellMerge", "tcPrChange",
]

TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
]

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]


def _clear_first_line(paragraph):
    """Kills inherited first-line indents (Word merges ind attributes across style inheritance)."""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        _insert_ordered(pPr, ind, PPR_ORDER)
    ind.attrib.pop(qn('w:hanging'), None)
    ind.set(qn('w:firstLineChars'), '0')
    ind.set(qn('w:firstLine'), '0')


class DocxPostProcessor:
    """
    Applies the final layout pass on the Pandoc output: table geometry and table grammar,
    caption recognition, block quote / code block rhythm and heading orphan control.
    """

    def __init__(self, docx_path: str, preset: str = "modern"):
        self.docx_path = docx_path
        self.preset = (preset or "modern").lower()
        if self.preset not in TABLE_DESIGN:
            self.preset = "modern"
        self.design = TABLE_DESIGN[self.preset]

    # ------------------------------------------------------------------ table helpers
    @staticmethod
    def _cell_margins(cell, pad: Tuple[int, int, int, int]):
        top, bottom, left, right = pad
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is None:
            tcMar = OxmlElement('w:tcMar')
            _insert_ordered(tcPr, tcMar, TCPR_ORDER)
        for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            node = tcMar.find(qn(f'w:{side}'))
            if node is None:
                node = OxmlElement(f'w:{side}')
                tcMar.append(node)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')

    @staticmethod
    def _cell_shading(cell, color: Optional[str]):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn('w:shd'))
        if color is None:
            if shd is not None:
                tcPr.remove(shd)
            return
        if shd is None:
            shd = OxmlElement('w:shd')
            _insert_ordered(tcPr, shd, TCPR_ORDER)
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), color)

    @staticmethod
    def _cell_borders(cell, top=None, bottom=None):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = tcPr.find(qn('w:tcBorders'))
        if borders is None:
            borders = OxmlElement('w:tcBorders')
            _insert_ordered(tcPr, borders, TCPR_ORDER)
        for side, spec in (("top", top), ("left", None), ("bottom", bottom), ("right", None)):
            node = borders.find(qn(f'w:{side}'))
            if node is None:
                node = OxmlElement(f'w:{side}')
                borders.append(node)
            if spec is None:
                node.set(qn('w:val'), 'none')
                node.set(qn('w:sz'), '0')
                node.set(qn('w:space'), '0')
                node.set(qn('w:color'), 'auto')
            else:
                val, sz, color = spec
                node.set(qn('w:val'), val)
                node.set(qn('w:sz'), str(sz))
                node.set(qn('w:space'), '0')
                node.set(qn('w:color'), color)

    @staticmethod
    def _cell_vcenter(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        vAlign = tcPr.find(qn('w:vAlign'))
        if vAlign is None:
            vAlign = OxmlElement('w:vAlign')
            _insert_ordered(tcPr, vAlign, TCPR_ORDER)
        vAlign.set(qn('w:val'), 'center')

    def _column_widths(self, table) -> List[int]:
        d = self.design
        size_pt = d["size"]
        pad = d["pad"][2] + d["pad"][3]
        mins, weights = [], []
        for ci in range(len(table.columns)):
            need_units = 4
            weight_units = 4
            for r_idx, cell in enumerate(table.columns[ci].cells):
                units = _cell_units(cell)
                # Header text is allowed to wrap onto a second line, body text is not.
                weight_units = max(weight_units, min(units, 11) if r_idx == 0 else units)
                need_units = max(need_units, min(units, 11) if r_idx == 0 else units)
            mins.append(int(need_units * size_pt * 10) + pad + 100)
            weights.append(min(max(weight_units, 6), 40))

        cap = int(CONTENT_W * 0.40)
        mins = [min(m, cap) for m in mins]
        total_min = sum(mins)
        if total_min < CONTENT_W:
            extra = CONTENT_W - total_min
            total_weight = sum(weights) or 1
            widths = [m + int(extra * w / total_weight) for m, w in zip(mins, weights)]
        else:
            # Short numeric columns must not be squeezed because one prose column is long:
            # give every column a floor of ~2 CJK chars and shrink the wide ones only.
            floor = int(5 * size_pt * 10) + pad
            floors = [min(m, floor) for m in mins]
            excess = total_min - CONTENT_W
            capacity = [m - f for m, f in zip(mins, floors)]
            total_capacity = sum(capacity) or 1
            widths = [m - int(excess * c / total_capacity) for m, c in zip(mins, capacity)]
            widths = [max(w, f) for w, f in zip(widths, floors)]

        diff = CONTENT_W - sum(widths)
        if diff:
            widest = widths.index(max(widths))
            widths[widest] += diff
        return widths

    def _style_table_text(self, cell, is_header: bool):
        d = self.design
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            _clear_first_line(p)
            for run in p.runs:
                run.font.name = d["latin"]
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is None:
                    rFonts = OxmlElement('w:rFonts')
                    rPr.insert(0, rFonts)
                rFonts.set(qn('w:eastAsia'), d["ea"])
                run.font.size = Pt(d["size"])
                if is_header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(d["header_color"])
                else:
                    run.font.color.rgb = RGBColor.from_string(d["color"])

    def enhance_tables(self, doc):
        d = self.design
        for table in doc.tables:
            tbl = table._tbl
            tblPr = tbl.tblPr

            borders = tblPr.find(qn('w:tblBorders'))
            if borders is not None:
                tblPr.remove(borders)
            layout = tblPr.find(qn('w:tblLayout'))
            if layout is None:
                layout = OxmlElement('w:tblLayout')
                _insert_ordered(tblPr, layout, TBLPR_ORDER)
            layout.set(qn('w:type'), 'fixed')
            tblW = tblPr.find(qn('w:tblW'))
            if tblW is None:
                tblW = OxmlElement('w:tblW')
                _insert_ordered(tblPr, tblW, TBLPR_ORDER)
            tblW.set(qn('w:type'), 'dxa')
            tblW.set(qn('w:w'), str(CONTENT_W))
            ind = tblPr.find(qn('w:tblInd'))
            if ind is None:
                ind = OxmlElement('w:tblInd')
                _insert_ordered(tblPr, ind, TBLPR_ORDER)
            ind.set(qn('w:type'), 'dxa')
            ind.set(qn('w:w'), '0')

            widths = self._column_widths(table)
            grid = tbl.tblGrid
            if grid is not None:
                for col, width in zip(grid.gridCol_lst, widths):
                    col.set(qn('w:w'), str(width))

            rows = table.rows
            last_idx = len(rows) - 1
            for r_idx, row in enumerate(rows):
                trPr = row._tr.get_or_add_trPr()
                if trPr.find(qn('w:cantSplit')) is None:
                    trPr.append(OxmlElement('w:cantSplit'))
                if r_idx == 0 and trPr.find(qn('w:tblHeader')) is None:
                    trPr.append(OxmlElement('w:tblHeader'))

                is_header = r_idx == 0
                is_last = r_idx == last_idx

                for c_idx, cell in enumerate(row.cells):
                    if c_idx < len(widths):
                        cell.width = Twips(widths[c_idx])
                    self._cell_vcenter(cell)
                    self._cell_margins(cell, d["pad"])

                    if is_header:
                        self._cell_shading(cell, d["header_fill"])
                        self._cell_borders(cell, top=d["header_top"], bottom=d["header_bottom"])
                    else:
                        fill = d["zebra"] if (d["zebra"] and r_idx % 2 == 1) else None
                        self._cell_shading(cell, fill)
                        bottom = d["last_bottom"] if is_last else d["row_line"]
                        self._cell_borders(cell, top=None, bottom=bottom)

                    self._style_table_text(cell, is_header)

    # ------------------------------------------------------------------ paragraph helpers
    def _style_caption(self, paragraph: Paragraph, kind: str):
        d = CAPTION_DESIGN[self.preset]
        style_name = 'Table Caption' if kind == 'table' else 'Image Caption'
        try:
            paragraph.style = paragraph.part.document.styles[style_name]
        except KeyError:
            pass
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        _clear_first_line(paragraph)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(3)
        if kind == 'table':
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.size = Pt(d["size"])
            run.font.bold = d["bold"]
            run.font.color.rgb = RGBColor.from_string(d["color"])
            run.font.italic = False

    def _style_code_block(self, paragraph: Paragraph):
        for run in paragraph.runs:
            rPr = run._element.find(qn('w:rPr'))
            if rPr is not None:
                rStyle = rPr.find(qn('w:rStyle'))
                if rStyle is not None and rStyle.get(qn('w:val')) == 'VerbatimChar':
                    rPr.remove(rStyle)
                shd = rPr.find(qn('w:shd'))
                if shd is not None:
                    rPr.remove(shd)
            run.font.name = "Consolas"
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.insert(0, rFonts)
            rFonts.set(qn('w:eastAsia'), "Consolas")
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.0

    @staticmethod
    def _style_quote_run(paragraphs: List[Paragraph], preset: str):
        if not paragraphs:
            return
        for p in paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
        paragraphs[0].paragraph_format.space_before = Pt(6)
        paragraphs[-1].paragraph_format.space_after = Pt(6)
        if preset == "academic":
            for p in paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)

    def enhance_body(self, doc):
        blocks = []
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:p'):
                blocks.append(Paragraph(child, doc))
            elif child.tag == qn('w:tbl'):
                blocks.append(Table(child, doc))

        quote_run: List[Paragraph] = []

        def flush_quotes():
            nonlocal quote_run
            self._style_quote_run(quote_run, self.preset)
            quote_run = []

        for i, block in enumerate(blocks):
            if isinstance(block, Table):
                flush_quotes()
                if i > 0 and isinstance(blocks[i - 1], Paragraph):
                    prev = blocks[i - 1]
                    text = prev.text.strip()
                    if text and len(text) <= 90 and TABLE_CAPTION_RE.match(text) \
                            and not text.endswith(("。", "：", ":")):
                        self._style_caption(prev, "table")
                continue

            style_name = (block.style.name or "").lower() if block.style is not None else ""

            if style_name in ("source code", "sourcecode"):
                flush_quotes()
                self._style_code_block(block)
                continue

            if "block" in style_name:
                quote_run.append(block)
            else:
                flush_quotes()

            if "heading" in style_name or "标题" in style_name:
                block.paragraph_format.keep_with_next = True

            has_drawing = bool(block._p.xpath('.//w:drawing | .//w:pict'))
            if has_drawing:
                block.alignment = WD_ALIGN_PARAGRAPH.CENTER
                block.paragraph_format.space_before = Pt(8)
                block.paragraph_format.space_after = Pt(4)
                if i + 1 < len(blocks) and isinstance(blocks[i + 1], Paragraph):
                    nxt = blocks[i + 1]
                    text = nxt.text.strip()
                    if text and len(text) <= 90 and FIG_CAPTION_RE.match(text):
                        self._style_caption(nxt, "figure")

        flush_quotes()

    def add_page_numbers(self, doc):
        size = 10.5 if self.preset == "academic" else 9
        color = "000000" if self.preset == "academic" else "8A94A6"
        for section in doc.sections:
            footer = section.footer
            if footer.is_linked_to_previous:
                footer.is_linked_to_previous = False
            paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            if paragraph.text.strip():
                continue
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            begin = OxmlElement('w:fldChar')
            begin.set(qn('w:fldCharType'), 'begin')
            instr = OxmlElement('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = ' PAGE '
            end = OxmlElement('w:fldChar')
            end.set(qn('w:fldCharType'), 'end')
            run._r.append(begin)
            run._r.append(instr)
            run._r.append(end)
            run.font.size = Pt(size)
            run.font.name = "Segoe UI" if self.preset != "academic" else "Times New Roman"
            run.font.color.rgb = RGBColor.from_string(color)

    def run(self):
        """Executes full post-processing."""
        try:
            doc = docx.Document(self.docx_path)
            self.enhance_tables(doc)
            self.enhance_body(doc)
            self.add_page_numbers(doc)
            doc.save(self.docx_path)
            logger.info(f"Post-processing complete for {os.path.basename(self.docx_path)}")
            return True
        except Exception as e:
            logger.warning(f"Post-processing failed for {os.path.basename(self.docx_path)}: {e}")
            return False
